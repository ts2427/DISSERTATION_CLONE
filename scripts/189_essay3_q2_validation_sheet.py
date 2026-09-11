"""
ESSAY 3 QUERY 2 — PART D1: BLIND VALIDATION SHEET (NEW)
=======================================================
Rows: the 50 calibration documents (outputs/rebuild/calibration_5_02/) + the 30-filing
validation draw fixed by scripts/187 (seed 20260911). Order randomised (seed 4242).
Each row shows the URL, the Item 5.02 section with the caption stripped (the frozen
classifier's parser, scripts/188), the reference date for "pre-announced" (the
reported_date of the earliest in-sample event whose window contains the filing), and
blank columns for the hand codes.

The classifier's answers go ONLY to validation_classifier_HIDDEN.csv (same sheet_id).
Nothing here prints them.
Outputs: outputs/essay3_q2/VALIDATION_SHEET.xlsx, validation_classifier_HIDDEN.csv
"""

import sys
import random
import importlib.util
from pathlib import Path
import pandas as pd

sys.stdout.reconfigure(encoding='utf-8', errors='replace')
OUT = Path('outputs/essay3_q2')
spec = importlib.util.spec_from_file_location('clf', 'scripts/188_essay3_q2_classifier.py')
clf = importlib.util.module_from_spec(spec)
spec.loader.exec_module(clf)

S = pd.read_csv(OUT / 'b_scope_filings.csv', dtype={'accession': str})
P = pd.read_csv(OUT / 'b_event_filing_pairs.csv', dtype={'accession': str})
ev = pd.read_csv('Data/processed/rebuild/CANONICAL_V3.csv', low_memory=False)
cal = pd.read_csv('outputs/rebuild/calibration_5_02/coding_sheet.csv', dtype={'accession': str})
val = pd.read_csv(OUT / 'validation_ids.csv', dtype=str)['accession'].tolist()

rows = []
for _, c in cal.iterrows():
    cik = int(c['local_file'].split('_')[1])
    rows.append(dict(accession=c['accession'], cik=cik, source='calibration',
                     path=f"outputs/rebuild/calibration_5_02/{c['local_file']}", filing_date=c['filing_date'],
                     org=c['org_name']))
for a in val:
    r = S[S['accession'] == a].iloc[0]
    org = ev.loc[ev['final_cik'] == r['cik'], 'org_name'].iloc[0]
    rows.append(dict(accession=a, cik=int(r['cik']), source='random30', path=r['local_file'],
                     filing_date=r['filing_date'], org=org))
V = pd.DataFrame(rows)
assert V['accession'].nunique() == 80, 'sheet must hold 80 distinct filings'
docs = {a: p for a, p in zip(S['accession'], S['primary_doc'])}

sheet, hidden = [], []
order = list(range(len(V)))
random.Random(4242).shuffle(order)
for k, i in enumerate(order, 1):
    r = V.iloc[i]
    t = clf.to_text(Path(r['path']).read_bytes())
    raw = clf.section_502(t)
    sec = clf.strip_caption(raw if raw else t[:20000])
    codes, persons, refs = clf.classify(sec)
    pp = P[P['accession'] == r['accession']]
    if len(pp):
        ref = pp.sort_values('reported_date').iloc[0]
        ref_rd, ref_bd = ref['reported_date'][:10], ref['breach_date']
    else:  # calibration document outside the Query 2 scope: use its calibration event
        e = ev[(ev['final_cik'] == r['cik']) & (ev['breach_date'] == cal.loc[cal['accession'] == r['accession'],
                                                                              'breach_date'].iloc[0])]
        ref_rd = str(e['reported_date'].iloc[0])[:10] if len(e) else ''
        ref_bd = e['breach_date'].iloc[0] if len(e) else ''
    doc = docs.get(r['accession']) or Path(r['path']).name.split('_', 3)[-1]
    url = f"https://www.sec.gov/Archives/edgar/data/{r['cik']}/{r['accession'].replace('-', '')}/{doc}"
    text = sec if len(sec) <= 32000 else sec[:32000] + ' [TRUNCATED — read the full filing at the URL]'
    sheet.append({'sheet_id': k, 'company': r['org'], 'filing_date': r['filing_date'], 'url': url,
                  'reference_date_for_pre_announced (event reported_date)': ref_rd,
                  'event breach_date': ref_bd, 'item_5_02_text (caption stripped)': text,
                  'EXEC officer departure (Y/N)': '', 'CEO departure (Y/N)': '',
                  'DIRECTOR-ONLY departure (Y/N)': '', 'person and role': '',
                  'pre-announced (Y/N/unclear)': '', 'notes': ''})
    deps = [p for p in persons if p['action'] == 'departure']
    hidden.append({'sheet_id': k, 'accession': r['accession'], 'source': r['source'],
                   'clf_exec_departure': 'Y' if codes['exec_departure'] else 'N',
                   'clf_ceo_departure': 'Y' if codes['ceo_departure'] else 'N',
                   'clf_director_only_departure': 'Y' if codes['director_departure'] else 'N',
                   'clf_person_role': '; '.join(f"{p['person']} ({p['role_class']}: {p['title']})" for p in deps),
                   'clf_pre_announced': (clf.pre_announced(codes, refs, pd.Timestamp(ref_rd)) if deps and ref_rd
                                         else ('N' if not deps else 'unclear')),
                   'clf_action': codes['action']})

instr = pd.DataFrame({'Coding rules (read first)': [
    'Code each filing from the Item 5.02 text shown (open the URL if the text is truncated or unclear).',
    'EXEC officer departure = Y if the filing reports that a person leaves (resigns, retires, is terminated, '
    'steps down, dies, or will leave on a stated date) a position as principal executive officer, president, '
    'principal financial officer, principal accounting officer, principal operating officer, or another '
    'named executive / executive officer. Appointments, elections and pay items alone are N.',
    'CEO departure = Y if the person leaving is the principal executive officer (CEO), including stepping '
    'down as CEO while staying on the board.',
    'DIRECTOR-ONLY departure = Y if a person who serves only as a director leaves the board (resigns, '
    'retires, will not stand for re-election, dies).',
    'person and role = name and title of each departing person, as the text states them.',
    'pre-announced = Y if the text references an announcement, agreement or notice about the departure '
    'dated BEFORE the reference date in column E; unclear if it says "as previously announced/disclosed" '
    'without a date; N otherwise (or if there is no departure).',
    'Do not open validation_classifier_HIDDEN.csv until every row is coded.']})
import re as _re
from openpyxl import load_workbook
from openpyxl.styles import Alignment
SH = pd.DataFrame(sheet).map(lambda v: _re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', ' ', v) if isinstance(v, str) else v)
xlsx = OUT / 'VALIDATION_SHEET.xlsx'
with pd.ExcelWriter(xlsx, engine='openpyxl') as xw:
    instr.to_excel(xw, sheet_name='README', index=False)
    SH.to_excel(xw, sheet_name='coding', index=False)
wb = load_workbook(xlsx)
ws = wb['coding']
for col, w in zip('ABCDEFGHIJKLM', [8, 26, 12, 40, 16, 14, 110, 14, 12, 14, 30, 14, 30]):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical='top')
wb['README'].column_dimensions['A'].width = 160
wb.save(xlsx)
pd.DataFrame(hidden).to_csv(OUT / 'validation_classifier_HIDDEN.csv', index=False)
print(f'VALIDATION_SHEET.xlsx: {len(sheet)} rows (calibration {int((V.source == "calibration").sum())}, '
      f'random30 {int((V.source == "random30").sum())}); hidden answers written separately.')
