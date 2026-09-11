"""
ESSAY 3 QUERY 2 — DIFFERENTIAL-RECALL AUDIT SHEET (NEW; v2 stays frozen — this measures, it does not revise)
===========================================================================================================
Why: across both validation rounds every miss of the final classifier v2 fell in a treated (carrier) filing
(round 2 out of sample: treated recall 1/3, control 2/2). If carrier filings word departures differently, the
misses bias the treated-control difference, not merely attenuate it. Tim's decision (2026-09-11): keep v2 frozen
and measure recall by treatment status on a blind audit stratified by treatment before estimating.

Population: B1-scope filings (in_spec_scope == 1) whose text was fetched, EXCLUDING every filing already coded or
read (calibration 50, round-1 random 30, round-2 30, development 40, Query-1 T-Mobile 35). Stratum = fcc_form499 of
the in-scope event the filing belongs to (earliest reported_date among the events whose B1 window contains it).
If a stratum has fewer than 40 eligible filings in the B1 scope, the extended scope (1,078) is used for that
stratum and the fact is logged.
Draw: 40 treated (seed 20260913) + 40 control (seed 20260914); row order randomised (seed 4444).
Same sheet format and README as rounds 1-2; the sheet shows NO treatment column. Classifier v2 answers and the
stratum go ONLY to validation_classifier_HIDDEN_AUDIT.csv.
Outputs: outputs/essay3_q2/audit_ids.csv, VALIDATION_SHEET_AUDIT.xlsx, validation_classifier_HIDDEN_AUDIT.csv
"""

import re
import sys
import random
import importlib.util
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment

sys.stdout.reconfigure(encoding='utf-8', errors='replace')
OUT = Path('outputs/essay3_q2')
spec = importlib.util.spec_from_file_location('clf2', 'scripts/195_essay3_q2_classifier_v2.py')
clf = importlib.util.module_from_spec(spec)
spec.loader.exec_module(clf)

S = pd.read_csv(OUT / 'b_scope_filings.csv', dtype={'accession': str})
P = pd.read_csv(OUT / 'b_event_filing_pairs.csv', dtype={'accession': str})
flog = pd.read_csv(OUT / 'b_fetch_log.csv', dtype={'accession': str})
ev = pd.read_csv('Data/processed/rebuild/CANONICAL_V3.csv', low_memory=False)
used = ({p.name.split('_')[2] for p in Path('outputs/rebuild/calibration_5_02').glob('*.htm')}
        | {p.name.split('_')[0] for p in Path('outputs/essay3_q1/tmobile_8k').glob('*.htm')}
        | set(pd.read_csv(OUT / 'validation_ids.csv', dtype=str)['accession'])
        | set(pd.read_csv(OUT / 'validation_ids_r2.csv', dtype=str)['accession'])
        | set(pd.read_csv(OUT / 'dev_ids.csv', dtype=str)['accession']))
ok = set(flog.loc[~flog['status'].str.startswith('FAILED'), 'accession'])


def stratum_frame(spec_only):
    pp = P[P['in_spec_scope'] == 1] if spec_only else P
    first = pp.sort_values('reported_date').drop_duplicates('accession')
    first = first[first['accession'].isin(ok) & ~first['accession'].isin(used)]
    return first[['accession', 'fcc_form499', 'reported_date', 'breach_date']]


spec_pool = stratum_frame(True)
full_pool = stratum_frame(False)
draw, log_lines = [], []
for g, seed in [(1, 20260913), (0, 20260914)]:
    pool = spec_pool[spec_pool['fcc_form499'] == g]
    scope = 'B1 scope'
    if len(pool) < 40:
        pool = full_pool[full_pool['fcc_form499'] == g]
        scope = 'extended scope (1,078)'
    ids = sorted(random.Random(seed).sample(sorted(pool['accession']), 40))
    log_lines.append(f"stratum {'treated' if g else 'control'}: eligible {len(pool)} in {scope}; drew 40 (seed {seed})")
    draw += [(a, g) for a in ids]
fp = OUT / 'audit_ids.csv'
D = pd.DataFrame(draw, columns=['accession', 'treated_stratum'])
if fp.exists():
    assert sorted(pd.read_csv(fp, dtype={'accession': str})['accession']) == sorted(D['accession']), 'audit draw must reproduce'
else:
    D.to_csv(fp, index=False)
for m in log_lines:
    print(m)

pairs = pd.concat([spec_pool, full_pool]).drop_duplicates('accession').set_index('accession')
order = list(range(len(D)))
random.Random(4444).shuffle(order)
sheet, hidden = [], []
for k, i in enumerate(order, 1):
    a, g = D.iloc[i]['accession'], int(D.iloc[i]['treated_stratum'])
    r = S[S['accession'] == a].iloc[0]
    t = clf.to_text(Path(r['local_file']).read_bytes())
    raw = clf.section_502(t)
    sec = clf.strip_caption(raw if raw else t[:20000])
    codes, persons, refs = clf.classify(sec)
    ref_rd, ref_bd = str(pairs.loc[a, 'reported_date'])[:10], pairs.loc[a, 'breach_date']
    org = ev.loc[ev['final_cik'] == r['cik'], 'org_name'].iloc[0]
    url = f"https://www.sec.gov/Archives/edgar/data/{r['cik']}/{a.replace('-', '')}/{r['primary_doc']}"
    text = sec if len(sec) <= 32000 else sec[:32000] + ' [TRUNCATED — read the full filing at the URL]'
    sheet.append({'sheet_id': k, 'company': org, 'filing_date': r['filing_date'], 'url': url,
                  'reference_date_for_pre_announced (event reported_date)': ref_rd,
                  'event breach_date': ref_bd, 'item_5_02_text (caption stripped)': text,
                  'EXEC officer departure (Y/N)': '', 'CEO departure (Y/N)': '',
                  'DIRECTOR-ONLY departure (Y/N)': '', 'person and role': '',
                  'pre-announced (Y/N/unclear)': '', 'notes': ''})
    deps = [p for p in persons if p['action'] == 'departure']
    hidden.append({'sheet_id': k, 'accession': a, 'cik': int(r['cik']), 'treated_stratum': g,
                   'clf_exec_departure': 'Y' if codes['exec_departure'] else 'N',
                   'clf_exec_departure_new': 'Y' if codes['exec_departure_new'] else 'N',
                   'clf_ceo_departure': 'Y' if codes['ceo_departure'] else 'N',
                   'clf_director_only_departure': 'Y' if codes['director_departure'] else 'N',
                   'clf_person_role': '; '.join(f"{p['person']} ({p['role_class']}: {p['title']})" for p in deps),
                   'clf_pre_announced': (clf.pre_announced(codes, refs, pd.Timestamp(ref_rd)) if deps else 'N'),
                   'clf_action': codes['action']})

instr = pd.read_excel(OUT / 'VALIDATION_SHEET.xlsx', sheet_name='README')
SH = pd.DataFrame(sheet).map(lambda v: re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', ' ', v) if isinstance(v, str) else v)
xlsx = OUT / 'VALIDATION_SHEET_AUDIT.xlsx'
with pd.ExcelWriter(xlsx, engine='openpyxl') as xw:
    instr.to_excel(xw, sheet_name='README', index=False)
    SH.to_excel(xw, sheet_name='coding', index=False)
wb = load_workbook(xlsx)
ws = wb['coding']
for col, w in zip('ABCDEFGHIJKLM', [8, 26, 12, 40, 16, 14, 110, 14, 12, 14, 30, 14, 30]):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical='top')
wb['README'].column_dimensions['A'].width = 160
wb.save(xlsx)
pd.DataFrame(hidden).to_csv(OUT / 'validation_classifier_HIDDEN_AUDIT.csv', index=False)
print(f'VALIDATION_SHEET_AUDIT.xlsx: {len(sheet)} rows (40 treated + 40 control strata, not shown on the sheet); '
      f'v2 answers only in validation_classifier_HIDDEN_AUDIT.csv')
