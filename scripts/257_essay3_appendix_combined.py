"""
ESSAY 3 APPENDIX OF RECORD - one combined document, tables renumbered to citation order
======================================================================================
scripts/245 writes Appendix Tables 1-11 and scripts/256 writes Tables 12-15 plus a note to
Table 5. This script reads those two committed .md outputs and writes ONE appendix in the order
the essay cites the tables. It changes LABELS, NOTES and DISPLAY FORMAT only. It estimates
nothing and does not touch 244, 245, 256 or anything they write.

  new  old  source
  1-7  1-7  245   unchanged numbers
   8   12   256   Randomization inference
   9    8   245   Pre-disclosure placebo
  10    9   245   Sensitivity analyses
  11   10   245   Cluster concentration
  12   11   245   T-Mobile events and executive departures
  13   15   256   T-Mobile filings quoted
   -   13   256   Chief executive counts      not carried: duplicates a column of Table 6
   -   14   256   Director-only counts        not carried: duplicates a column of Table 6

What changes, and nothing else:
  - three note sentences (Tables 2 and 5), each an exact substitution that must match once;
  - Table 5 note: one sentence on which classifier the audit scored, replacing the 256 note;
  - Table 6 note: T-Mobile's share of treated director-only events, computed here from
    e_analysis_sample.csv and c2_departure_events.csv and asserted against Table 6;
  - Table 8: coefficients to two decimals without a plus sign (from the committed
    randomization file, asserted equal to the Table 7 and Table 9 coefficients), window
    labels as in the 245 tables;
  - Table 11: cluster names from Table 2 where the parent CIK is listed there;
  - Table 12 note: which classifier-coded departures Panel B does not list, each part of the
    sentence asserted against c2_departure_events.csv and c2_person_rows.csv;
  - Table 13: "PLACEBO" in sentence case;
  - no source line, script name, repository path or commit hash, and no numbered hypothesis
    label, anywhere in the document (asserted on the written file).
Every other cell is read back from the written file and compared with the 245 / 256 cell.

Outputs (outputs/essay3_appendix/): ESSAY3_APPENDIX.md, ESSAY3_APPENDIX.docx,
TABLE_NUMBER_CROSSWALK.csv, APPENDIX_RENUMBER_CHECK.csv
"""

import hashlib
import json
import re
import sys
from pathlib import Path

import pandas as pd

sys.stdout.reconfigure(encoding='utf-8', errors='replace')
APP = Path('outputs/essay3_appendix')
V4 = Path('outputs/essay3_v4')
SRC245 = APP / 'ESSAY3_APPENDIX_TABLES.md'
SRC256 = APP / 'ESSAY3_APPENDIX_SUPPLEMENT.md'

# new number -> (source script, old number)
ORDER = {1: (245, 1), 2: (245, 2), 3: (245, 3), 4: (245, 4), 5: (245, 5), 6: (245, 6), 7: (245, 7),
         8: (256, 12), 9: (245, 8), 10: (245, 9), 11: (245, 10), 12: (245, 11), 13: (256, 15)}
NEW = {v: k for k, v in ORDER.items()}
# 256 tables that are not carried into the appendix of record, and why
DROPPED = {(256, 13): 'not in the appendix of record: duplicates the chief executive column of Table 6',
           (256, 14): 'not in the appendix of record: duplicates the director-only column of Table 6'}
# the file under each old number that holds the table's data (filenames keep their old numbers)
INTERNAL = {(245, n): f'outputs/essay3_q4/table{n:02d}.csv' for n in range(1, 12)}
INTERNAL.update({(256, 12): 'outputs/essay3_appendix/supp_table12_randomization_inference.csv',
                 (256, 13): 'outputs/essay3_appendix/supp_table13_ceo_departure_counts.csv',
                 (256, 14): 'outputs/essay3_appendix/supp_table14_director_departure_counts.csv',
                 (256, 15): 'outputs/essay3_appendix/supp_table15_tmobile_departures.csv'})

# (source, old number, old text, new text): each must match exactly once in that table's note
NOTE_FIXES = [
    (245, 2, 'they are one corporate family only in the entity count (12)',
     'they are one parent entity only in the entity count (12)'),
    (245, 2, 'the Boost Mobile divestiture that the date-conditional rule turns on',
     "the date DISH's acquisition of Boost Mobile closed, on which the date-conditional rule turns"),
    (245, 5, 'An empty cell means the field had no reference positives and no classifier positives, '
             'so the statistic is undefined.',
     'An empty cell means the statistic is undefined: recall when the field had no reference '
     'positives, and precision and kappa when it also had no classifier positives.'),
]
# cross-references to another table by number, found in the notes: (source, old number of the
# note's table, text as written, text after renumbering). The scan below stops on any
# reference that is not listed here.
XREFS = [
    (256, 12, 're-estimates the specification of Table 7', 're-estimates the specification of Table 7'),
]
# the appendix of record carries no source line, script name or repository path: each of these
# is removed from the note it sits in, and must match exactly once
REMOVALS = [
    (256, 12, ' Source: `outputs/defense_supplement/e3_randomization_inference.csv` (`scripts/250`).'),
    (256, 15, ' Sources: `outputs/essay3_v4/g6_case_table.csv` (`scripts/228`); '
              '`outputs/essay3_q4/tmobile_502_text.md` (verbatim filing text).'),
]
# replaces the 256 note to Table 5; 256 itself checks the claim (same classifier functions)
AUDIT_SENTENCE = ('The stratified audit scored the same classifier version that the analysis sample uses; '
                  'the later relinking changed the sample, not the classifier.')
HYP = re.compile(r'\bH\d+[a-z]?\b')
BANNED = ('outputs/', 'scripts/', '.csv', '.md', '`', 'Source:', 'Sources:')


def stop(msg):
    print('STOP: ' + msg)
    sys.exit(1)


def cells(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def parse(path):
    """-> (front matter lines, {old number: table}, {old number: text of a 'Note to Table N'})."""
    lines = path.read_text(encoding='utf-8').split('\n')
    heads = [i for i, ln in enumerate(lines) if re.fullmatch(r'\*\*Table \d+\*\*', ln)]
    front, tables, extra = lines[:heads[0]], {}, {}
    for i, ln in enumerate(front):
        m = re.match(r'## Note to Table (\d+)', ln)
        if m:
            extra[int(m.group(1))] = next(x for x in front[i + 1:] if x.strip())
    for a, b in zip(heads, heads[1:] + [len(lines)]):
        num = int(re.search(r'\d+', lines[a]).group(0))
        body = [x for x in lines[a + 1:b] if x.strip()]
        title, rest = body[0], body[1:]
        if not (title.startswith('*') and title.endswith('*')):
            stop(f'{path.name} Table {num}: no title line')
        blocks, label, rows, note = [], '', [], None
        for x in rest:
            if x.startswith('|'):
                rows.append(x)
                continue
            if rows:
                blocks.append((label, rows))
                label, rows = '', []
            if x.startswith('*Note.*'):
                note = x
            elif re.fullmatch(r'\*\*.+\*\*', x):
                label = x.strip('*')
            else:
                stop(f'{path.name} Table {num}: unparsed line {x[:60]!r}')
        if rows or note is None or not blocks:
            stop(f'{path.name} Table {num}: expected panels then one note')
        for lab, rws in blocks:
            w = len(cells(rws[0]))
            if any(len(cells(r)) != w for r in rws):
                stop(f'{path.name} Table {num} {lab}: ragged table')
        tables[num] = dict(title=title.strip('*'), blocks=blocks, note=note)
    return front, tables, extra


def once(text, old, new, where):
    n = text.count(old)
    if n != 1:
        stop(f'{where}: {old!r} matches {n} times, expected exactly 1')
    return text.replace(old, new)


def row(c):
    return '| ' + ' | '.join(c) + ' |'


NUM = re.compile(r'[−+-]?(?:\d[\d,]*(?:\.\d+)?|\.\d+)')


def numeric(tab):
    """Every number in every body cell, in reading order."""
    return [[NUM.findall(c) for r in rows[2:] for c in cells(r)] for _, rows in tab['blocks']]


def digest(obj):
    return hashlib.sha256(json.dumps(obj, ensure_ascii=False).encode('utf-8')).hexdigest()


front245, T245, _ = parse(SRC245)
_, T256, EXTRA = parse(SRC256)
if sorted(T245) != list(range(1, 12)) or sorted(T256) != [12, 13, 14, 15] or sorted(EXTRA) != [5]:
    stop(f'unexpected source tables: 245 {sorted(T245)}, 256 {sorted(T256)}, notes {sorted(EXTRA)}')
SRC = {(245, n): t for n, t in T245.items()}
SRC.update({(256, n): t for n, t in T256.items()})
if sorted(SRC) != sorted(list(NEW) + list(DROPPED)):
    stop('ORDER and DROPPED do not cover every source table exactly once')
BEFORE = {k: (digest(numeric(t)), digest(t['blocks'])) for k, t in SRC.items()}

# ------------------------------------------------------------------ notes: fixes and cross-references
NOTES = {k: SRC[k]['note'] for k in NEW}
for src, old, a, b in NOTE_FIXES:
    NOTES[(src, old)] = once(NOTES[(src, old)], a, b, f'note fix, {src} Table {old}')
handled = {}
for src, old, a, b in XREFS:
    NOTES[(src, old)] = once(NOTES[(src, old)], a, b, f'cross-reference, {src} Table {old}')
    handled[(src, old)] = handled.get((src, old), 0) + len(re.findall(r'Tables? \d+', a))
for k in NEW:
    t = SRC[k]
    found = len(re.findall(r'Tables? \d+', t['note'])) + sum(
        len(re.findall(r'Tables? \d+', ' '.join([lab] + rows))) for lab, rows in t['blocks'])
    if found != handled.get(k, 0):
        stop(f'{k[0]} Table {k[1]}: {found} table references, {handled.get(k, 0)} handled')
for src, old, a in REMOVALS:
    NOTES[(src, old)] = once(NOTES[(src, old)], a, '', f'source line, {src} Table {old}')
if 'are identical to those in `scripts/195`' not in EXTRA[5]:
    stop('the 256 note to Table 5 no longer states that the audited classifier is the one in use')
NOTES[(245, 5)] += ' ' + AUDIT_SENTENCE
HYP_REPLACED = 0
for k in NOTES:
    NOTES[k], n = HYP.subn('the hypothesis test', NOTES[k])
    HYP_REPLACED += n

# ------------------------------------------------------------------ Table 6 note: T-Mobile and director-only
# A director-only event is a treated event with at least one director-only departure first
# disclosed in (t0, t0 + w]. The event flags are rebuilt here from the departure file, checked
# against the analysis sample, and the totals are checked against Table 6.
TM = 1283699
smp = pd.read_csv(V4 / 'e_analysis_sample.csv', low_memory=False)
smp = smp[smp['in_analysis_sample'] == 1]
trt = smp[smp['fcc_form499'] == 1]
tmo = trt[trt['final_cik'] == TM]
if (len(smp), len(trt), len(tmo)) != (405, 109, 26) or set(tmo['outcome_cik']) != {TM}:
    stop('analysis sample is not 405 / 109 treated / 26 T-Mobile')
dep = pd.read_csv(V4 / 'c2_departure_events.csv', dtype={'first_accession': str})
dep = dep[(dep['cik'] == TM) & (dep['grp'] == 'director') & (dep['director_only'] == 1)]
dep_date = pd.to_datetime(dep['first_date'])
t6 = [cells(r) for r in T245[6]['blocks'][0][1]]
jd = t6[0].index('Director-only departure')
DIR = {}
for w in (30, 90, 180):
    col, files, n_tm = f'director_departure_{w}_rd', set(), 0
    for _, e in tmo.iterrows():
        days = (dep_date - pd.Timestamp(e['reported_date'])).dt.days
        hit = dep[(days > 0) & (days <= w)]
        if int(len(hit) > 0) != int(e[col]):
            stop(f'T-Mobile event {e["breach_date"]}: rebuilt {w}-day director-only flag differs from the sample')
        n_tm += int(len(hit) > 0)
        files |= set(hit['first_accession'])
    n_all = int(trt[col].sum())
    six = [x[jd] for x in t6[2:] if x[0] == str(w) and x[1] == 'Treated']
    if six != [f'{n_all} ({100 * n_all / len(trt):.1f}%)']:
        stop(f'treated director-only count at {w} days ({n_all}) is not the Table 6 cell {six}')
    DIR[w] = dict(tm=n_tm, all=n_all, files=len(files))
rest30, rest_n = DIR[30]['all'] - DIR[30]['tm'], len(trt) - len(tmo)
if DIR != {30: dict(tm=5, all=8, files=3), 90: dict(tm=12, all=18, files=8),
           180: dict(tm=13, all=21, files=8)} or (rest30, rest_n) != (3, 83):
    stop(f'director-only figures moved: {DIR}, {rest30} of {rest_n}')
DIR_SENTENCE = ('T-Mobile events are ' + ', '.join(
    f'{DIR[w]["tm"]} of the {DIR[w]["all"]} treated events with a director-only departure at {w} days '
    f'({100 * DIR[w]["tm"] / DIR[w]["all"]:.1f}%)' for w in (30, 90, 180))
    + f'; those T-Mobile events trace to {DIR[30]["files"]}, {DIR[90]["files"]} and {DIR[180]["files"]} '
    'distinct first-disclosure filings, because one filing falls in the window of several events. '
    f'At 30 days, {rest30} of the {rest_n} treated events outside T-Mobile have a director-only departure.')
NOTES[(245, 6)] += ' ' + DIR_SENTENCE

# ------------------------------------------------------------------ Table 12 note: codes Panel B does not list
# Panel B of the T-Mobile table is a fixed list of six persons. The sentence below says which
# classifier-coded executive departures it leaves out; each part of it is checked here.
PANEL_B_SENTENCE = (
    'Panel B lists the four outcome-window departures and the 2019 succession filing. Three other '
    'classifier-coded departures fall only in placebo windows and are not listed: two coded from a filing '
    'dated April 30, 2018, which the filing text does not describe as departures, and one for David Carey '
    'from a filing dated February 19, 2020. Only the placebo flag of the 2018-08-20 event depends on an '
    'unlisted code.')
pb = [cells(r) for r in T245[11]['blocks'][1][1]]
iNm, iFd, iAc, iWp = (pb[0].index(x) for x in ('Name', 'Filing date', 'Accession', 'Window placement'))
listed = {(c[iNm].split()[-1].lower(), c[iAc]) for c in pb[2:]}
succession = {(c[iFd], c[iAc]) for c in pb[2:] if c[iWp] == 'Placebo window'}
if [c[iWp] for c in pb[2:]].count('Outcome window') != 4 or succession != {('2019-11-18', '0001193125-19-294093')}:
    stop('Panel B is not four outcome-window departures and the 2019 succession filing')
exd = pd.read_csv(V4 / 'c2_departure_events.csv', dtype={'first_accession': str})
exd = exd[(exd['cik'] == TM) & (exd['grp'] == 'exec')]
exd_date = pd.to_datetime(exd['first_date'])
unlisted, in_outcome, rests_on_unlisted = set(), set(), []
for _, e in tmo.iterrows():
    days = (exd_date - pd.Timestamp(e['reported_date'])).dt.days
    plc = exd[(days > -180) & (days <= 0)]
    keys = [(k, a, f) for k, a, f in zip(plc['pkey'], plc['first_accession'], plc['first_date'])]
    if int(len(plc) > 0) != int(e['placebo_exec_departure_rd']):
        stop(f'T-Mobile event {e["breach_date"]}: rebuilt placebo flag differs from the sample')
    out = [(k, a, f) for k, a, f in keys if (k, a) not in listed]
    unlisted |= set(out)
    if keys and len(out) == len(keys):
        rests_on_unlisted.append(str(e['breach_date'])[:10])
    owin = exd[(days > 0) & (days <= 180)]
    in_outcome |= {(k, a) for k, a in zip(owin['pkey'], owin['first_accession'])}
if unlisted != {('legere', '0001104659-18-028086', '2018-04-30'), ('sievert', '0001104659-18-028086', '2018-04-30'),
                ('carey', '0001193125-20-041926', '2020-02-19')}:
    stop(f'unlisted placebo-window codes are not the three expected: {sorted(unlisted)}')
if {(k, a) for k, a, _ in unlisted} & in_outcome or rests_on_unlisted != ['2018-08-20']:
    stop(f'an unlisted code is in an outcome window, or the flags resting on one are {rests_on_unlisted}')
if in_outcome - listed:
    stop(f'an outcome-window departure is missing from Panel B: {sorted(in_outcome - listed)}')
# the two 2018 codes come from a succession-of-title sentence and a severance clause
verbs = pd.read_csv(V4 / 'c2_person_rows.csv', dtype=str)
verbs = verbs[(verbs['accession'] == '0001104659-18-028086') & (verbs['action'] == 'departure')]
if sorted(verbs['verb']) != ['succeed/replace', 'termination of the employment']:
    stop(f'the 2018 filing is no longer coded from the two expected phrases: {sorted(verbs["verb"])}')
NOTES[(245, 11)] += ' ' + PANEL_B_SENTENCE

# ------------------------------------------------------------------ Table 8 (old 12): display format
# Two decimals and no plus sign, as in the 245 tables. The two-decimal coefficients come from
# the committed randomization file and must equal the coefficients printed in Tables 7 and 9.
ri = pd.read_csv('outputs/defense_supplement/e3_randomization_inference.csv')
t7c, prev = {}, ''
for r in T245[7]['blocks'][0][1][2:]:  # the window is printed once per group of rows
    c = cells(r)
    prev = c[0] or prev
    if c[1] == 'CV3':
        t7c[prev] = c[2]
t9c = [cells(r)[1] for r in T245[8]['blocks'][0][1][2:] if cells(r)[0] == 'CV3']
if sorted(t7c) != ['180', '30', '90'] or len(t9c) != 1:
    stop('could not read the CV3 coefficients from the primary and placebo tables')
riA = SRC[(256, 12)]['blocks'][0]
h = cells(riA[1][0])
if h[0] != 'Outcome window' or h[3] != 'Coefficient (pp)' or h[4] != 'CV3 t' or len(ri) != len(riA[1]) - 2:
    stop('randomization table Panel A is not laid out as expected')
h[0] = 'Window (days)'
newRI = [row(h), riA[1][1]]
for r, (_, x) in zip(riA[1][2:], ri.iterrows()):
    c = cells(r)
    two = f'{100 * x["coef_obs"]:.2f}'.replace('-', '−')
    if f'{100 * x["coef_obs"]:+.1f}'.replace('-', '−') != c[3] or f'{x["cv3_t_obs"]:+.2f}'.replace('-', '−') != c[4]:
        stop(f'randomization table row {c[0]} / {c[1]} is not the row of the committed file')
    want = t7c[str(int(x['window']))] if x['outcome'].startswith('exec_departure') else t9c[0]
    if two != want:
        stop(f'randomization coefficient {two} is not the coefficient {want} of the estimates table')
    c[0], c[3], c[4] = re.sub(r'^(\d+) days$', r'\1', c[0]), two, c[4].lstrip('+')
    newRI.append(row(c))
RI_BLOCKS = [(riA[0], newRI)] + SRC[(256, 12)]['blocks'][1:]

# ------------------------------------------------------------------ Table 13 (old 15): sentence case
q = SRC[(256, 15)]['blocks'][0]
Q_ROWS = q[1][:]
hits = [i for i, r in enumerate(Q_ROWS) if 'PLACEBO' in r]
if len(hits) != 1 or Q_ROWS[hits[0]].count('PLACEBO (before notification)') != 1:
    stop('expected exactly one PLACEBO cell in the T-Mobile filings table')
Q_ROWS[hits[0]] = Q_ROWS[hits[0]].replace('PLACEBO (before notification)', 'Placebo (before notification)')

# ------------------------------------------------------------------ cluster names (new Table 11)
# One name per parent CIK, taken from Table 2 Panel A where the CIK is listed there; any other
# cluster keeps the name it has. Panel A of the cluster table prints names without a CIK, so
# those are looked up through the committed variance-share file the names came from.
t2 = [cells(r) for r in T245[2]['blocks'][0][1][2:]]
T2NAME = {int(r[0].replace(',', '')): r[1] for r in t2 if r[0] != 'Total'}
if len(T2NAME) != 13:
    stop(f'Table 2 Panel A lists {len(T2NAME)} parent CIKs, expected 13')
vs = pd.read_csv(V4 / 'f1_cv3_variance_shares.csv')[['final_cik', 'name']].drop_duplicates()
if vs['name'].duplicated().any() or vs['final_cik'].duplicated().any():
    stop('f1_cv3_variance_shares.csv: name and parent CIK are not one to one')
CIK_OF = dict(zip(vs['name'], vs['final_cik'].astype(int)))
RENAMED = []


def rename(cik, name, where):
    new = T2NAME.get(cik, name)
    if new != name:
        RENAMED.append((where, cik, name, new))
    return new


cc = SRC[(245, 10)]
(labA, rowsA), (labB, rowsB), panC = cc['blocks']
hA, hB = cells(rowsA[0]), cells(rowsB[0])
iR, iC, iN, iW = hA.index('Reversing clusters (β without)'), hB.index('CIK'), hB.index('Cluster'), 0
newA = rowsA[:2]
for r in rowsA[2:]:
    c = cells(r)
    parts = []
    for item in c[iR].split('; '):
        m = re.fullmatch(r'(.+) \(([−\d.]+)\)', item)
        if not m or m.group(1) not in CIK_OF:
            stop(f'cluster table Panel A: cannot resolve {item!r} to a parent CIK')
        parts.append(f'{rename(CIK_OF[m.group(1)], m.group(1), f"Panel A, {c[iW]} days")} ({m.group(2)})')
    c[iR] = '; '.join(parts)
    newA.append(row(c))
newB = rowsB[:2]
for r in rowsB[2:]:
    c = cells(r)
    cik = int(c[iC].replace(',', ''))
    if CIK_OF.get(c[iN]) != cik:
        stop(f'cluster table Panel B: {c[iN]!r} is not the name on file for CIK {cik}')
    c[iN] = rename(cik, c[iN], f'Panel B, {c[iW]} days')
    newB.append(row(c))

OUT_T = {k: dict(SRC[k], note=NOTES[k]) for k in NEW}
OUT_T[(245, 10)]['blocks'] = [(labA, newA), (labB, newB), panC]
OUT_T[(256, 12)]['blocks'] = RI_BLOCKS
OUT_T[(256, 15)]['blocks'] = [(q[0], Q_ROWS)]

# ------------------------------------------------------------------ write .md
pre = [x for x in front245 if x.strip() and not x.startswith('# ')]
if len(pre) != 2 or not pre[1].startswith('Conventions:'):
    stop('245 front matter is not the preamble and the conventions paragraph')
PREAMBLE, CONVENTIONS = pre
RECORD = ('This is the Essay 3 appendix of record. The two build files it is assembled from are internal '
          'sources and keep their old table numbers; the table number crosswalk maps each old number to the '
          'number used here.')
L = ['# Essay 3 Appendix', '', RECORD, '', PREAMBLE, '', CONVENTIONS, '']
for new in sorted(ORDER):
    t = OUT_T[ORDER[new]]
    L += [f'**Table {new}**', '', f'*{t["title"]}*', '']
    for lab, rows in t['blocks']:
        if lab:
            L += [f'**{lab}**', '']
        L += rows + ['']
    L += [t['note'], '']
(APP / 'ESSAY3_APPENDIX.md').write_text('\n'.join(L) + '\n', encoding='utf-8')

# ------------------------------------------------------------------ crosswalk
cw = pd.DataFrame([dict(new_number=NEW.get((src, old), ''), old_number=old, source_script=f'scripts/{src}',
                        source_document=(SRC245 if src == 245 else SRC256).as_posix(),
                        internal_id=INTERNAL[(src, old)], title=SRC[(src, old)]['title'],
                        status=DROPPED.get((src, old), 'in the appendix of record'))
                   for src, old in sorted(SRC, key=lambda k: (NEW.get(k, 99), k[1]))])
for f in cw['internal_id']:
    if not Path(f).exists():
        stop(f'crosswalk names a file that does not exist: {f}')
cw.to_csv(APP / 'TABLE_NUMBER_CROSSWALK.csv', index=False, lineterminator='\n')

# ------------------------------------------------------------------ checks on the written file
TXT = (APP / 'ESSAY3_APPENDIX.md').read_text(encoding='utf-8')
_, BACK, _ = parse(APP / 'ESSAY3_APPENDIX.md')
if sorted(BACK) != list(range(1, len(ORDER) + 1)):
    stop(f'combined appendix tables are {sorted(BACK)}, expected 1-{len(ORDER)} in order')
leak = [w for w in BANNED if w in TXT] + HYP.findall(TXT) + re.findall(r'\b(?=[0-9a-f]*[a-f])(?=[a-f]*\d)[0-9a-f]{7}\b', TXT)
if leak or 'PLACEBO' in TXT:
    stop(f'appendix of record still holds a path, source line, hash, hypothesis label or PLACEBO: {leak}')
if any(re.match(r'\+|\d+ days$', c) for r in BACK[NEW[(256, 12)]]['blocks'][0][1][2:] for c in cells(r)):
    stop('a randomization table cell still carries a plus sign or a "days" window label')
# cells allowed to differ from the 245 / 256 cell: (old text, new text)
allowed = {(245, 10): {(o, n) for _, _, o, n in RENAMED} | {
               (cells(r1)[iR], cells(r2)[iR]) for r1, r2 in zip(rowsA[2:], newA[2:]) if r1 != r2},
           (256, 12): {(x, y) for r1, r2 in zip(riA[1], newRI) for x, y in zip(cells(r1), cells(r2)) if x != y},
           (256, 15): {('PLACEBO (before notification)', 'Placebo (before notification)')}}
chk = []
for new, key in sorted(ORDER.items()):
    b, s, o = BACK[new], SRC[key], OUT_T[key]
    if b['blocks'] != o['blocks'] or b['note'] != o['note']:
        stop(f'Table {new}: the written file does not read back as built')
    num_b, txt_b = digest(numeric(b)), digest(b['blocks'])
    diff = [(x, y) for (_, rs), (_, rb) in zip(s['blocks'], b['blocks'])
            for r1, r2 in zip(rs, rb) for x, y in zip(cells(r1), cells(r2)) if x != y]
    if [(lab, len(rows)) for lab, rows in b['blocks']] != [(lab, len(rows)) for lab, rows in s['blocks']]:
        stop(f'Table {new} (old {key[1]}): panels or row counts differ from the {key[0]} output')
    if set(diff) - allowed.get(key, set()):
        stop(f'Table {new} (old {key[1]}): a cell differs from the {key[0]} output: {sorted(set(diff))[:3]}')
    if key != (256, 12) and num_b != BEFORE[key][0]:
        stop(f'Table {new} (old {key[1]}): a numeric cell differs from the {key[0]} output')
    if b['title'] != s['title']:
        stop(f'Table {new}: title changed')
    chk.append(dict(new_number=new, old_number=key[1], source_script=f'scripts/{key[0]}',
                    body_cells=sum(len(cells(r)) for _, rs in b['blocks'] for r in rs[2:]),
                    numbers=sum(len(c) for blk in numeric(b) for c in blk),
                    sha256_numeric_source=BEFORE[key][0], sha256_numeric_combined=num_b,
                    numeric_identical=num_b == BEFORE[key][0],
                    body_text_identical=txt_b == BEFORE[key][1], cells_changed=len(diff),
                    note_changed=b['note'] != s['note']))
pd.DataFrame(chk).to_csv(APP / 'APPENDIX_RENUMBER_CHECK.csv', index=False, lineterminator='\n')

# the two 256 tables that are not carried must still equal the matching column of Table 6
# Panel A, by window and group: that is the reason they are not carried
t6b = [cells(r) for r in BACK[6]['blocks'][0][1]]
for old, col in ((13, 'Chief executive departure'), (14, 'Director-only departure')):
    j = t6b[0].index(col)
    for r in (cells(x) for x in SRC[(256, old)]['blocks'][0][1][2:]):
        for grp, cnt, rate in (('Treated', r[1], r[3]), ('Control', r[2], r[4])):
            six = [x[j] for x in t6b[2:] if f'{x[0]} days' == r[0] and x[1] == grp]
            if six != [f'{cnt} ({rate})']:
                stop(f'256 Table {old} {r[0]} {grp}: {cnt} ({rate}) but Table 6 has {six}')

# ------------------------------------------------------------------ write .docx
from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402

doc = Document()
st = doc.styles['Normal']
st.font.name, st.font.size = 'Times New Roman', Pt(12)
st.paragraph_format.space_after, st.paragraph_format.line_spacing = Pt(0), 1.0
for s in doc.sections:
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)


def para(runs, size=12, align=None, sa=6):
    p = doc.add_paragraph()
    for text, bold, italic in runs:
        r = p.add_run(text)
        r.font.name, r.font.size, r.bold, r.italic = 'Times New Roman', Pt(size), bold, italic
    p.paragraph_format.space_after, p.paragraph_format.line_spacing = Pt(sa), 1.0
    if align is not None:
        p.alignment = align
    return p


para([('Essay 3 Appendix', True, False)], size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
for x in (RECORD, PREAMBLE, CONVENTIONS):
    para([(x, False, False)], size=11)
WIDE = {NEW[(245, 7)], NEW[(245, 9)], NEW[(245, 10)], NEW[(256, 15)]}
for new in sorted(ORDER):
    t = OUT_T[ORDER[new]]
    if new in WIDE:
        s = doc.add_section()
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = s.page_height, s.page_width
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)
    else:
        doc.add_page_break()
    para([(f'Table {new}', True, False)], sa=0)
    para([(t['title'], False, True)], sa=8)
    for lab, rows in t['blocks']:
        if lab:
            para([(lab, True, False)], size=11, sa=4)
        tb = doc.add_table(rows=0, cols=len(cells(rows[0])))
        tb.style = 'Table Grid'
        for k, r in enumerate([rows[0]] + rows[2:]):
            for cell, v in zip(tb.add_row().cells, cells(r)):
                run = cell.paragraphs[0].add_run(v)
                run.font.size, run.font.name, run.bold = Pt(9), 'Times New Roman', k == 0
                cell.paragraphs[0].paragraph_format.line_spacing = 1.0
        para([('', False, False)], sa=4)
    para([('Note. ', False, True), (t['note'][len('*Note.*'):].strip(), False, False)], size=9)
doc.save(str(APP / 'ESSAY3_APPENDIX.docx'))

# read the saved .docx back: same numbers in order, same cells as the .md
d = Document(str(APP / 'ESSAY3_APPENDIX.docx'))
heads = [p.text for p in d.paragraphs if re.fullmatch(r'Table \d+', p.text)]
if heads != [f'Table {n}' for n in range(1, len(ORDER) + 1)]:
    stop(f'.docx table headings are {heads}')
md_cells = [c for n in range(1, len(ORDER) + 1) for _, rows in BACK[n]['blocks']
            for r in [rows[0]] + rows[2:] for c in cells(r)]
dx_cells = [c.text for tb in d.tables for row_ in tb.rows for c in row_.cells]
if md_cells != dx_cells:
    stop('.docx table cells differ from the .md')

same = [c['new_number'] for c in chk if c['numeric_identical']]
print(f'wrote {APP / "ESSAY3_APPENDIX.md"} and .docx: {len(ORDER)} tables in citation order')
print(f'note fixes applied exactly once: {len(NOTE_FIXES)}; cross-references checked: {len(XREFS)}; '
      f'source lines removed: {len(REMOVALS)}; hypothesis labels replaced: {HYP_REPLACED}')
print(f'numeric cells identical to the 245/256 outputs in Tables {same}; Table 8 coefficients reformatted '
      'to two decimals and checked against Tables 7 and 9')
print('256 Tables 13 and 14 (not carried) equal the chief executive and director-only columns of Table 6')
print('Table 6 note: ' + DIR_SENTENCE)
print(f'cluster names changed in Table 11: {len(RENAMED)}')
for where, cik, a, b in RENAMED:
    print(f'  {where}: CIK {cik}: {a} -> {b}')
