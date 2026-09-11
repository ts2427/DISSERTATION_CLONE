"""
ESSAY 3 QUERY 2 — PART D3 ROUND 2: FRESH BLIND VALIDATION SHEET FOR CLASSIFIER v2 (NEW)
=====================================================================================
Revision rule: the revised classifier (scripts/195, committed before this draw) is re-validated on a fresh
random draw of 30 filings, coded blind.
Population: the same as round 1 — B1-scope filings (b_scope_filings.csv, in_spec_scope == 1) whose text was
fetched — EXCLUDING the 80 round-1 sheet filings (50 calibration + 30 random), the 40 development filings
(dev_ids.csv), and the 35 T-Mobile documents read in Query 1. Draw: 30, seed 20260912. Order randomised
(seed 4343). Same sheet format as round 1 (scripts/189). Classifier v2 answers go ONLY to
validation_classifier_HIDDEN_R2.csv.
Outputs: outputs/essay3_q2/validation_ids_r2.csv, VALIDATION_SHEET_R2.xlsx, validation_classifier_HIDDEN_R2.csv
"""

import re
import sys
import random
import importlib.util
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment

sys.stdout.reconfigure(encoding='utf-8', errors='replace')
OUT = Path('outputs/essay3_q2')
spec = importlib.util.spec_from_file_location('clf2', 'scripts/195_essay3_q2_classifier_v2.py')
clf = importlib.util.module_from_spec(spec)
spec.loader.exec_module(clf)

S = pd.read_csv(OUT / 'b_scope_filings.csv', dtype={'accession': str})
P = pd.read_csv(OUT / 'b_event_filing_pairs.csv', dtype={'accession': str})
flog = pd.read_csv(OUT / 'b_fetch_log.csv', dtype={'accession': str})
ev = pd.read_csv('Data/processed/rebuild/CANONICAL_V3.csv', low_memory=False)
cal = {p.name.split('_')[2] for p in Path('outputs/rebuild/calibration_5_02').glob('*.htm')}
tmo = {p.name.split('_')[0] for p in Path('outputs/essay3_q1/tmobile_8k').glob('*.htm')}
r1 = set(pd.read_csv(OUT / 'validation_ids.csv', dtype=str)['accession'])
dev = set(pd.read_csv(OUT / 'dev_ids.csv', dtype=str)['accession'])
ok = set(flog.loc[~flog['status'].str.startswith('FAILED'), 'accession'])
pool = sorted(a for a in S.loc[S['in_spec_scope'] == 1, 'accession']
              if a in ok and a not in cal and a not in r1 and a not in dev and a not in tmo)
draw = sorted(random.Random(20260912).sample(pool, 30))
fp = OUT / 'validation_ids_r2.csv'
if fp.exists():
    assert sorted(pd.read_csv(fp, dtype=str)['accession']) == draw, 'round-2 draw must reproduce exactly'
else:
    pd.DataFrame({'accession': draw}).to_csv(fp, index=False)
print(f'Round-2 pool {len(pool)} (B1 scope minus calibration 50, round-1 30, dev 40, Query-1 T-Mobile 35); '
      f'drew 30 with seed 20260912')

sheet, hidden = [], []
order = list(range(30))
random.Random(4343).shuffle(order)
for k, i in enumerate(order, 1):
    a = draw[i]
    r = S[S['accession'] == a].iloc[0]
    t = clf.to_text(Path(r['local_file']).read_bytes())
    raw = clf.section_502(t)
    sec = clf.strip_caption(raw if raw else t[:20000])
    codes, persons, refs = clf.classify(sec)
    pp = P[(P['accession'] == a) & (P['in_spec_scope'] == 1)].sort_values('reported_date')
    ref_rd, ref_bd = str(pp['reported_date'].iloc[0])[:10], pp['breach_date'].iloc[0]
    org = ev.loc[ev['final_cik'] == r['cik'], 'org_name'].iloc[0]
    url = f"https://www.sec.gov/Archives/edgar/data/{r['cik']}/{a.replace('-', '')}/{r['primary_doc']}"
    text = sec if len(sec) <= 32000 else sec[:32000] + ' [TRUNCATED — read the full filing at the URL]'
    sheet.append({'sheet_id': k, 'company': org, 'filing_date': r['filing_date'], 'url': url,
                  'reference_date_for_pre_announced (event reported_date)': ref_rd,
                  'event breach_date': ref_bd, 'item_5_02_text (caption stripped)': text,
                  'EXEC officer departure (Y/N)': '', 'CEO departure (Y/N)': '',
                  'DIRECTOR-ONLY departure (Y/N)': '', 'person and role': '',
                  'pre-announced (Y/N/unclear)': '', 'notes': ''})
    deps = [p for p in persons if p['action'] == 'departure']
    hidden.append({'sheet_id': k, 'accession': a, 'source': 'random30_r2',
                   'clf_exec_departure': 'Y' if codes['exec_departure'] else 'N',
                   'clf_exec_departure_new': 'Y' if codes['exec_departure_new'] else 'N',
                   'clf_ceo_departure': 'Y' if codes['ceo_departure'] else 'N',
                   'clf_ceo_departure_new': 'Y' if codes['ceo_departure_new'] else 'N',
                   'clf_director_only_departure': 'Y' if codes['director_departure'] else 'N',
                   'clf_person_role': '; '.join(f"{p['person']} ({p['role_class']}: {p['title']})" for p in deps),
                   'clf_pre_announced': (clf.pre_announced(codes, refs, pd.Timestamp(ref_rd)) if deps else 'N'),
                   'clf_restated_mention': codes['restated_mention'], 'clf_vacancy_mention': codes['vacancy_mention'],
                   'clf_implied_succession': codes['implied_succession'], 'clf_action': codes['action']})

instr = pd.read_excel(OUT / 'VALIDATION_SHEET.xlsx', sheet_name='README')  # identical rules to round 1
SH = pd.DataFrame(sheet).map(lambda v: re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', ' ', v) if isinstance(v, str) else v)
xlsx = OUT / 'VALIDATION_SHEET_R2.xlsx'
with pd.ExcelWriter(xlsx, engine='openpyxl') as xw:
    instr.to_excel(xw, sheet_name='README', index=False)
    SH.to_excel(xw, sheet_name='coding', index=False)
wb = load_workbook(xlsx)
ws = wb['coding']
for col, w in zip('ABCDEFGHIJKLM', [8, 26, 12, 40, 16, 14, 110, 14, 12, 14, 30, 14, 30]):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical='top')
wb['README'].column_dimensions['A'].width = 160
wb.save(xlsx)
pd.DataFrame(hidden).to_csv(OUT / 'validation_classifier_HIDDEN_R2.csv', index=False)
print(f'VALIDATION_SHEET_R2.xlsx: {len(sheet)} rows; v2 answers written only to validation_classifier_HIDDEN_R2.csv')
